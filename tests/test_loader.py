from dataclasses import FrozenInstanceError
from pathlib import Path

import docx
import pymupdf
import pytest

from contract_rag.loader import (
    SUPPORTED_EXTENSIONS,
    LoaderError,
    LoaderErrorCode,
    NoTextLayerError,
    PageText,
    UnsupportedFormatError,
    join_hyphenation,
    load_document,
    load_docx,
    load_pdf,
    load_txt,
    strip_boilerplate,
)

CORPUS_DIR = Path(__file__).parents[1] / "corpus_raw" / "contracts"
CORPUS_FILES = sorted(CORPUS_DIR.glob("contract_*.pdf"))
FIXTURES_DIR = Path(__file__).parent / "fixtures"


def _create_pdf(path: Path, page_texts: list[str]) -> None:
    with pymupdf.open() as document:
        for text in page_texts:
            page = document.new_page()
            if text:
                page.insert_text((72, 72), text)
        document.save(path)


def _create_docx(
    path: Path,
    paragraphs: list[str],
    table_rows: list[tuple[str, str]] | None = None,
) -> None:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=2)
        for row_index, (left, right) in enumerate(table_rows):
            table.rows[row_index].cells[0].text = left
            table.rows[row_index].cells[1].text = right
    document.save(path)


def test_page_text_is_immutable_and_requires_positive_page_number() -> None:
    page = PageText(page_number=1, text="body", source_file="sample.pdf")

    with pytest.raises(FrozenInstanceError):
        page.text = "changed"  # type: ignore[misc]
    with pytest.raises(ValueError, match="greater than or equal to 1"):
        PageText(page_number=0, text="body", source_file="sample.pdf")


def test_load_pdf_returns_one_based_page_objects(tmp_path: Path) -> None:
    pdf_path = tmp_path / "three-pages.pdf"
    _create_pdf(pdf_path, ["Alpha", "Beta", "Gamma"])

    pages = load_pdf(pdf_path)

    assert [page.page_number for page in pages] == [1, 2, 3]
    assert [page.text for page in pages] == ["Alpha", "Beta", "Gamma"]
    assert all(page.source_file == str(pdf_path) for page in pages)


def test_strip_boilerplate_normalizes_digits_and_whitespace_at_exact_position() -> None:
    pages = [
        PageText(1, "  Agreement   1\nBody alpha\nPage 1", "sample.pdf"),
        PageText(2, "Agreement 22\nBody beta\nPage   2", "sample.pdf"),
        PageText(3, "Third heading\nShared\nThird footer", "sample.pdf"),
        PageText(4, "Fourth heading\nBody delta\nFourth footer", "sample.pdf"),
    ]

    cleaned = strip_boilerplate(pages)

    assert [page.text for page in cleaned] == [
        "Body alpha",
        "Body beta",
        "Third heading\nShared\nThird footer",
        "Fourth heading\nBody delta\nFourth footer",
    ]
    assert [page.text for page in pages] == [
        "  Agreement   1\nBody alpha\nPage 1",
        "Agreement 22\nBody beta\nPage   2",
        "Third heading\nShared\nThird footer",
        "Fourth heading\nBody delta\nFourth footer",
    ]


def test_strip_boilerplate_requires_two_occurrences_and_one_source() -> None:
    single_page = [PageText(1, "\nHeading\nBody\nFooter\n", "sample.pdf")]

    assert strip_boilerplate(single_page) == [PageText(1, "Heading\nBody\nFooter", "sample.pdf")]
    assert strip_boilerplate([]) == []
    with pytest.raises(ValueError, match="same source_file"):
        strip_boilerplate(
            [
                PageText(1, "First", "first.pdf"),
                PageText(1, "Second", "second.pdf"),
            ]
        )


def test_strip_boilerplate_keeps_matches_from_different_positions() -> None:
    pages = [
        PageText(1, "Shared\nBody alpha\nFooter alpha", "sample.pdf"),
        PageText(2, "Header beta\nShared\nFooter beta", "sample.pdf"),
        PageText(3, "Header gamma\nBody gamma\nFooter gamma", "sample.pdf"),
        PageText(4, "Header delta\nBody delta\nFooter delta", "sample.pdf"),
    ]

    assert strip_boilerplate(pages) == pages


def test_strip_boilerplate_counts_empty_pages_in_threshold() -> None:
    pages = [
        PageText(1, "Marker\nBody alpha", "sample.pdf"),
        PageText(2, "Marker\nBody beta", "sample.pdf"),
        PageText(3, "", "sample.pdf"),
        PageText(4, "Header delta\nBody delta", "sample.pdf"),
        PageText(5, "Header epsilon\nBody epsilon", "sample.pdf"),
    ]

    assert strip_boilerplate(pages) == pages


@pytest.mark.parametrize(
    ("source", "expected"),
    [
        (
            "...policies and prac-\ntices have pushed...",
            "...policies and practices have pushed...",
        ),
        ("...Auch andere EU-\nStaaten, wie...", "...Auch andere EU-Staaten, wie..."),
        ("обяза-\n    тельства сторон", "обязательства сторон"),
        ("келіс-\n    ім талаптары", "келісім талаптары"),
        ("multi\u2010\nlingual", "multilingual"),
        ("EU-\r\n\tStaaten", "EU-Staaten"),
        ("end-\n\nnew paragraph", "end-\n\nnew paragraph"),
        ("section-\n2", "section-\n2"),
    ],
)
def test_join_hyphenation(source: str, expected: str) -> None:
    assert join_hyphenation(source) == expected


def test_load_pdf_raises_no_text_layer_error_for_blank_pdf(tmp_path: Path) -> None:
    pdf_path = tmp_path / "blank.pdf"
    _create_pdf(pdf_path, ["", ""])

    with pytest.raises(NoTextLayerError) as exception_info:
        load_pdf(pdf_path)

    assert exception_info.value.code is LoaderErrorCode.NO_TEXT_LAYER
    assert exception_info.value.source_file == str(pdf_path)


def test_load_pdf_reports_expected_file_errors(tmp_path: Path) -> None:
    missing_path = tmp_path / "missing.pdf"
    with pytest.raises(LoaderError) as missing_error:
        load_pdf(missing_path)
    assert missing_error.value.code is LoaderErrorCode.FILE_NOT_FOUND

    with pytest.raises(LoaderError) as directory_error:
        load_pdf(tmp_path)
    assert directory_error.value.code is LoaderErrorCode.PATH_IS_DIRECTORY

    invalid_path = tmp_path / "invalid.pdf"
    invalid_path.write_text("not a PDF", encoding="utf-8")
    with pytest.raises(LoaderError) as invalid_error:
        load_pdf(invalid_path)
    assert invalid_error.value.code is LoaderErrorCode.INVALID_PDF


def test_load_pdf_rejects_password_protected_pdf(tmp_path: Path) -> None:
    pdf_path = tmp_path / "protected.pdf"
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "Protected text")
        document.save(
            pdf_path,
            encryption=pymupdf.PDF_ENCRYPT_AES_256,
            owner_pw="owner-secret",
            user_pw="user-secret",
        )

    with pytest.raises(LoaderError) as exception_info:
        load_pdf(pdf_path)

    assert exception_info.value.code is LoaderErrorCode.PASSWORD_PROTECTED


@pytest.mark.parametrize("pdf_path", CORPUS_FILES, ids=lambda path: path.name)
def test_load_pdf_handles_every_contract_in_corpus(pdf_path: Path) -> None:
    with pymupdf.open(pdf_path) as document:
        expected_page_count = document.page_count

    pages = load_pdf(pdf_path)

    assert len(pages) == expected_page_count
    assert [page.page_number for page in pages] == list(range(1, expected_page_count + 1))
    assert all(page.source_file == str(pdf_path) for page in pages)
    assert any(page.text for page in pages)


def test_page_text_accepts_a_missing_page_number() -> None:
    page = PageText(page_number=None, text="body", source_file="sample.docx")

    assert page.page_number is None
    with pytest.raises(ValueError, match="greater than or equal to 1"):
        PageText(page_number=0, text="body", source_file="sample.docx")


def test_load_document_reads_docx_as_one_pageless_page(tmp_path: Path) -> None:
    docx_path = tmp_path / "agreement.docx"
    _create_docx(docx_path, ["1.1 First clause.", "1.2 Second clause."])

    pages = load_document(docx_path)

    assert len(pages) == 1
    assert pages[0].page_number is None
    assert pages[0].text == "1.1 First clause.\n1.2 Second clause."
    assert pages[0].source_file == str(docx_path)


def test_load_document_reads_txt_as_one_pageless_page(tmp_path: Path) -> None:
    txt_path = tmp_path / "agreement.txt"
    txt_path.write_text("1.1 First clause.\n1.2 Second clause.\n", encoding="utf-8")

    pages = load_document(txt_path)

    assert len(pages) == 1
    assert pages[0].page_number is None
    assert pages[0].text == "1.1 First clause.\n1.2 Second clause."
    assert pages[0].source_file == str(txt_path)


def test_load_document_is_identical_to_load_pdf_for_pdf(tmp_path: Path) -> None:
    pdf_path = tmp_path / "three-pages.pdf"
    _create_pdf(pdf_path, ["Alpha", "Beta", "Gamma"])

    assert load_document(pdf_path) == load_pdf(pdf_path)


@pytest.mark.parametrize("filename", ["notes.rtf", "archive.pdf.zip", "no-extension"])
def test_load_document_rejects_unsupported_formats(tmp_path: Path, filename: str) -> None:
    unsupported_path = tmp_path / filename
    unsupported_path.write_text("content", encoding="utf-8")

    with pytest.raises(UnsupportedFormatError) as exception_info:
        load_document(unsupported_path)

    assert exception_info.value.code is LoaderErrorCode.UNSUPPORTED_FORMAT
    assert exception_info.value.source_file == str(unsupported_path)
    assert SUPPORTED_EXTENSIONS == {".pdf", ".docx", ".txt"}


def test_load_document_matches_extensions_case_insensitively(tmp_path: Path) -> None:
    txt_path = tmp_path / "AGREEMENT.TXT"
    txt_path.write_text("1.1 Clause.", encoding="utf-8")

    assert load_document(txt_path)[0].text == "1.1 Clause."


def test_load_docx_keeps_table_text_in_document_order(tmp_path: Path) -> None:
    docx_path = tmp_path / "with-table.docx"
    _create_docx(
        docx_path,
        ["3. РЕКВИЗИТЫ СТОРОН"],
        table_rows=[("Исполнитель", "Заказчик"), ("БИН 123", "БИН 456")],
    )

    text = load_docx(docx_path)[0].text

    assert text.splitlines() == [
        "3. РЕКВИЗИТЫ СТОРОН",
        "Исполнитель\tЗаказчик",
        "БИН 123\tБИН 456",
    ]


def test_load_docx_joins_hyphenation_like_the_pdf_path(tmp_path: Path) -> None:
    docx_path = tmp_path / "hyphenated.docx"
    _create_docx(docx_path, ["1.1 в согласованном ассорти-", "менте товара."])

    assert load_docx(docx_path)[0].text == "1.1 в согласованном ассортименте товара."


def test_load_docx_reports_expected_file_errors(tmp_path: Path) -> None:
    with pytest.raises(LoaderError) as missing_error:
        load_docx(tmp_path / "missing.docx")
    assert missing_error.value.code is LoaderErrorCode.FILE_NOT_FOUND

    with pytest.raises(LoaderError) as directory_error:
        load_docx(tmp_path)
    assert directory_error.value.code is LoaderErrorCode.PATH_IS_DIRECTORY

    invalid_path = tmp_path / "invalid.docx"
    invalid_path.write_text("not a DOCX", encoding="utf-8")
    with pytest.raises(LoaderError) as invalid_error:
        load_docx(invalid_path)
    assert invalid_error.value.code is LoaderErrorCode.INVALID_DOCX


def test_load_txt_reports_expected_file_errors(tmp_path: Path) -> None:
    with pytest.raises(LoaderError) as missing_error:
        load_txt(tmp_path / "missing.txt")
    assert missing_error.value.code is LoaderErrorCode.FILE_NOT_FOUND

    with pytest.raises(LoaderError) as directory_error:
        load_txt(tmp_path)
    assert directory_error.value.code is LoaderErrorCode.PATH_IS_DIRECTORY


def test_load_txt_rejects_non_utf8_bytes(tmp_path: Path) -> None:
    txt_path = tmp_path / "cp1251.txt"
    txt_path.write_bytes("Договор поставки".encode("cp1251"))

    with pytest.raises(LoaderError) as exception_info:
        load_txt(txt_path)

    assert exception_info.value.code is LoaderErrorCode.INVALID_ENCODING


def test_load_txt_strips_the_utf8_byte_order_mark(tmp_path: Path) -> None:
    txt_path = tmp_path / "bom.txt"
    txt_path.write_text("1.1 Clause.", encoding="utf-8-sig")

    assert load_txt(txt_path)[0].text == "1.1 Clause."


def test_empty_docx_and_txt_raise_no_text_layer(tmp_path: Path) -> None:
    empty_docx = tmp_path / "empty.docx"
    _create_docx(empty_docx, ["", "   "])
    with pytest.raises(NoTextLayerError) as docx_error:
        load_docx(empty_docx)
    assert docx_error.value.code is LoaderErrorCode.NO_TEXT_LAYER

    empty_txt = tmp_path / "empty.txt"
    empty_txt.write_text("\n  \n", encoding="utf-8")
    with pytest.raises(NoTextLayerError) as txt_error:
        load_txt(empty_txt)
    assert txt_error.value.code is LoaderErrorCode.NO_TEXT_LAYER


@pytest.mark.parametrize(
    "fixture_name",
    ["sample_contract.docx", "sample_contract.txt"],
)
def test_committed_fixtures_load_as_one_pageless_page(fixture_name: str) -> None:
    fixture_path = FIXTURES_DIR / fixture_name

    pages = load_document(fixture_path)

    assert len(pages) == 1
    assert pages[0].page_number is None
    assert pages[0].source_file == str(fixture_path)
    assert "2.2 Оплата производится" in pages[0].text


def test_committed_docx_fixture_keeps_bank_details_from_its_table() -> None:
    text = load_document(FIXTURES_DIR / "sample_contract.docx")[0].text

    assert "IBAN KZ11 1111 1111 1111" in text
    assert "БИН 210987654321" in text


def test_committed_txt_fixture_joins_hyphenated_words() -> None:
    text = load_document(FIXTURES_DIR / "sample_contract.txt")[0].text

    assert "ассортименте" in text
    assert "ассорти-" not in text
